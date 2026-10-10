import re
import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=120, bottom=120, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="CCCCCC"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="6" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="single" w:sz="6" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="single" w:sz="6" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def build_word_report(md_path, docx_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    doc = Document()

    # Thiet lap le tieu chuan 1 inch (2.54 cm)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Thong nhat 1 phong chu Times New Roman toan bo tai lieu
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(13)
    font.color.rgb = RGBColor(0x11, 0x11, 0x11)
    style.paragraph_format.line_spacing = 1.25
    style.paragraph_format.space_after = Pt(4)

    i = 0
    in_code_block = False
    code_lines = []
    
    in_table = False
    table_rows = []

    def flush_table():
        nonlocal in_table, table_rows
        if not table_rows:
            in_table = False
            return
        
        num_cols = max(len(r) for r in table_rows)
        table = doc.add_table(rows=len(table_rows), cols=num_cols)
        table.autofit = True
        set_table_borders(table)

        for r_idx, row_data in enumerate(table_rows):
            is_header = (r_idx == 0)
            for c_idx in range(num_cols):
                cell = table.cell(r_idx, c_idx)
                text = row_data[c_idx] if c_idx < len(row_data) else ""
                cell.text = text.strip()
                set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
                
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.line_spacing = 1.15
                    for run in p.runs:
                        run.font.name = 'Times New Roman'
                        run.font.size = Pt(11)
                        if is_header:
                            run.font.bold = True
                            run.font.color.rgb = RGBColor(0x00, 0x20, 0x60)

                if is_header:
                    set_cell_background(cell, "EAEFF5")
                elif r_idx % 2 == 1:
                    set_cell_background(cell, "F9FAFB")
        
        doc.add_paragraph()
        table_rows = []
        in_table = False

    is_cover_page = True

    while i < len(lines):
        line = lines[i].rstrip('\r\n')

        # Khoi ma nguon (Code box)
        if line.startswith("```"):
            if not in_code_block:
                in_code_block = True
                code_lines = []
            else:
                in_code_block = False
                code_text = "\n".join(code_lines)
                tbl = doc.add_table(rows=1, cols=1)
                cell = tbl.cell(0, 0)
                set_cell_background(cell, "F5F6F8")
                set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
                cell.text = code_text
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.line_spacing = 1.15
                    for run in p.runs:
                        run.font.name = 'Consolas'
                        run.font.size = Pt(9.5)
                        run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
                doc.add_paragraph()
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Bang bieu Markdown
        if "|" in line and not line.startswith("```"):
            stripped = line.strip()
            if stripped.startswith("|") and stripped.endswith("|"):
                parts = [p.strip() for p in stripped.split("|")[1:-1]]
                if all(re.match(r'^:?-+:?$', p) for p in parts if p):
                    i += 1
                    continue
                in_table = True
                table_rows.append(parts)
                i += 1
                continue
        
        if in_table:
            flush_table()

        clean = line.strip()

        # Cac dau moc ngat trang chuyen muc chinh
        page_break_titles = [
            "# LỜI MỞ ĐẦU VÀ TRI ÂN",
            "# CAM ĐOAN KẾT QUẢ ĐỒ ÁN",
            "# BẢNG ĐÁNH GIÁ CỦA GIẢNG VIÊN",
            "# MỤC LỤC CHI TIẾT",
            "# CHƯƠNG 1: KHẢO SÁT HIỆN TRẠNG VÀ NHẬN DIỆN KHIẾM KHUYẾT HỆ THỐNG V1.0",
            "# CHƯƠNG 2: ĐỀ XUẤT KIẾN TRÚC MỚI VÀ LUẬN GIẢI MẪU THIẾT KẾ CHUẨN SYLLABUS",
            "# CHƯƠNG 3: THIẾT KẾ KIẾN TRÚC TỔNG THỂ HỆ THỐNG V2.0 VÀ LUỒNG TƯƠNG TÁC",
            "# CHƯƠNG 4: ĐÁNH GIÁ CHUYÊN SÂU TUÂN THỦ NGUYÊN TẮC SOLID VÀ NGUYÊN LÝ GOF",
            "# CHƯƠNG 5: KỊCH BẢN KIỂM THỬ THỰC NGHIỆM VÀ ĐÁNH GIÁ VẬN HÀNH",
            "# CHƯƠNG 6: KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN",
            "# TÀI LIỆU THAM KHẢO"
        ]

        if clean in page_break_titles:
            doc.add_page_break()
            is_cover_page = False

        # Tieu de cac cap (Dung chung phong chu Times New Roman)
        if line.startswith("# "):
            h_text = line[2:].strip()
            h = doc.add_heading(level=1)
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(6)
            run = h.add_run(h_text)
            run.font.name = 'Times New Roman'
            run.font.bold = True
            run.font.size = Pt(16)
            run.font.color.rgb = RGBColor(0x00, 0x20, 0x60)
            if is_cover_page or "LỜI MỞ ĐẦU" in h_text or "CAM ĐOAN" in h_text or "BẢNG ĐÁNH GIÁ" in h_text or "MỤC LỤC" in h_text or "TÀI LIỆU THAM KHẢO" in h_text:
                h.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue
        elif line.startswith("## "):
            h = doc.add_heading(level=2)
            h.paragraph_format.space_before = Pt(12)
            h.paragraph_format.space_after = Pt(4)
            run = h.add_run(line[3:].strip())
            run.font.name = 'Times New Roman'
            run.font.bold = True
            run.font.size = Pt(14)
            run.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
            i += 1
            continue
        elif line.startswith("### "):
            h = doc.add_heading(level=3)
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(3)
            run = h.add_run(line[4:].strip())
            run.font.name = 'Times New Roman'
            run.font.bold = True
            run.font.size = Pt(13)
            run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
            i += 1
            continue
        elif line.startswith("#### "):
            h = doc.add_heading(level=4)
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(2)
            run = h.add_run(line[5:].strip())
            run.font.name = 'Times New Roman'
            run.font.bold = True
            run.font.size = Pt(12)
            run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
            i += 1
            continue

        if line.strip() in ["---", "***", "___"]:
            i += 1
            continue

        if not clean:
            i += 1
            continue

        p = doc.add_paragraph()
        if is_cover_page and ("THÀNH PHỐ HỒ CHÍ MINH" in clean or "BÁO CÁO TIỂU LUẬN" in clean or "MÃ MÔN HỌC" in clean):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif clean.startswith("* ") or clean.startswith("- "):
            p.paragraph_format.left_indent = Inches(0.25)
            clean = clean[2:]
        elif re.match(r'^\d+\.\s', clean):
            m = re.match(r'^(\d+\.\s)', clean)
            p.paragraph_format.left_indent = Inches(0.25)
            clean = clean[len(m.group(1)):]

        # Khong dung mui ten, phong chu nhat quan
        tokens = re.split(r'(\*\*.*?\*\*|`.*?`|\*.*?\*)', clean)
        for tok in tokens:
            if tok.startswith("**") and tok.endswith("**"):
                r = p.add_run(tok[2:-2])
                r.font.name = 'Times New Roman'
                r.font.bold = True
            elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
                r = p.add_run(tok[1:-1])
                r.font.name = 'Times New Roman'
                r.font.italic = True
            elif tok.startswith("`") and tok.endswith("`"):
                r = p.add_run(tok[1:-1])
                r.font.name = 'Times New Roman'
                r.font.bold = True
                r.font.color.rgb = RGBColor(0x80, 0x00, 0x00)
            else:
                r = p.add_run(tok)
                r.font.name = 'Times New Roman'

        i += 1

    if in_table:
        flush_table()

    try:
        doc.save(docx_path)
        print("Exported successfully to BAO_CAO_TIET_LUAN_GIUA_KY.docx")
    except PermissionError:
        alt_path = os.path.join(os.path.dirname(docx_path), "BAO_CAO_TIET_LUAN_GIUA_KY_v2.docx")
        doc.save(alt_path)
        print("Do file dang duoc mo trong Microsoft Word, da luu thanh cong thanh: BAO_CAO_TIET_LUAN_GIUA_KY_v2.docx")

if __name__ == "__main__":
    script_dir = os.path.dirname(os.path.abspath(__file__))
    md_file = os.path.join(script_dir, "BAO_CAO_TIET_LUAN_GIUA_KY.md")
    docx_file = os.path.join(script_dir, "BAO_CAO_TIET_LUAN_GIUA_KY.docx")
    build_word_report(md_file, docx_file)
